"""Document library: built-in templates (no organization names, valid Word files), official versions with
major/minor numbering and variants, approval by a second admin, superseding, replaceable templates, who sees what,
write-once files, upload safety, tenant isolation and the audit log."""
import hashlib
import io
import re
import zipfile

import pytest
import yaml
from fastapi.testclient import TestClient

from grcpal.portal import doclib
from grcpal.portal.app import create_app
from grcpal.portal.service import Forbidden, Invalid
from grcpal.portal.store import SqliteStore
from test_portal import ADMIN, ADMIN2, AUDITOR, CONTRIB, HR, PDF, as_, csrf, make_portal

DOCX = doclib.builtin_template("p-05")[1]  # any real Word file will do as an "uploaded" document


@pytest.fixture()
def portal(tmp_path):
    return make_portal(tmp_path)


@pytest.fixture()
def client(portal, monkeypatch):
    monkeypatch.setenv("GRCPAL_AUTH_MODE", "dev")
    monkeypatch.setenv("GRCPAL_ENV", "dev")
    monkeypatch.delenv("GRCPAL_PUBLIC_URL", raising=False)
    return TestClient(create_app(portal), base_url="https://grc.example.com")


def actions(portal):
    return [r["action"] for r in portal.log.records()]


def official(portal, did="p-05", **kw):
    args = dict(variant="", change="major", comment="")
    args.update(kw)
    return portal.upload_official(ADMIN, did, "access-control.docx", DOCX, args["variant"], args["change"], args["comment"])


# ------------------------------------------------------------------ the built-in catalog and templates
def test_catalog_is_complete_and_consistent():
    cat = doclib.builtin_catalog()
    docs, folders = cat["docs"], {f["id"] for f in cat["folders"]}
    assert len(docs) >= 50 and len({d["number"] for d in docs.values()}) == len(docs)
    for d in docs.values():
        assert d["folder"] in folders and d["title"] and doclib.builtin_template(d["id"]), d["id"]
        assert not d.get("parent") or d["parent"] in docs, d["id"]
    kinds = {d["folder"]: 0 for d in docs.values()}
    for d in docs.values():
        kinds[d["folder"]] += 1
    assert kinds["tier1"] == 1 and kinds["tier2"] == 13 and kinds["tier3"] == 13 and kinds["tier4"] == 20 and kinds["soc2"] == 6


def test_templates_are_valid_word_files_without_org_names():
    from docx import Document

    for did in doclib.builtin_catalog()["docs"]:
        name, data = doclib.builtin_template(did)
        assert name == f"{did}.docx" and data[:4] == b"PK\x03\x04"
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        assert "[Company Name]" in text, did
        assert not re.search(r"\bHSP\b|Azure|Entra|Intune|SharePoint|VP of Security", text), did


def test_full_policies_have_testable_statements():
    from docx import Document

    for did in ["isp"] + [f"p-{n:02d}" for n in range(1, 14)]:
        doc = Document(io.BytesIO(doclib.builtin_template(did)[1]))
        heads = [t.rows[0].cells[2].text for t in doc.tables if len(t.columns) == 3]
        assert heads and all(h == "How it is tested" for h in heads), did
        statements = sum(len(t.rows) - 1 for t in doc.tables if len(t.columns) == 3)
        assert statements >= 6, did


def test_catalog_file_matches_the_sources():
    src = yaml.safe_load(open("tools/doctemplates/registry.yaml"))
    assert {d["id"] for d in src["docs"]} == set(doclib.builtin_catalog()["docs"])


def test_control_ids_in_catalog_exist(portal):
    known = {c["id"] for c in portal.catalog.controls}
    for d in doclib.builtin_catalog()["docs"].values():
        for part in re.findall(r"[A-Z]+-\d+", d.get("controls", "")):
            assert part in known, (d["id"], part)


# ------------------------------------------------------------------ library list and permissions
def test_admin_and_auditor_see_everything_staff_only_approved(portal):
    lib = portal.library(ADMIN)
    assert lib["total"] == len(doclib.builtin_catalog()["docs"]) and {f["id"] for f in lib["folders"]} >= {"tier1", "tier2"}
    assert portal.library(AUDITOR)["total"] == lib["total"]
    assert portal.library(CONTRIB)["total"] == 0
    with pytest.raises(Invalid):
        portal.library_doc(CONTRIB, "p-05")
    v = official(portal)
    portal.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved")
    assert [d["id"] for d in portal.library(CONTRIB)["docs"]] == ["p-05"]
    assert portal.library_doc(CONTRIB, "p-05")["current"][0]["label"] == "1.0"
    only = portal.library(ADMIN, folder="tier2", q="access")
    assert [d["id"] for d in only["docs"]] == ["p-05"]


def test_only_admins_write(portal):
    for who in (CONTRIB, AUDITOR):
        with pytest.raises(Forbidden):
            portal.upload_official(who, "p-05", "a.docx", DOCX, "", "major", "")
        with pytest.raises(Forbidden):
            portal.upload_template(who, "p-05", "a.docx", DOCX, "")
        with pytest.raises(Forbidden):
            portal.edit_doc(who, "p-05", {"title": "Changed title"})
        with pytest.raises(Forbidden):
            portal.add_doc(who, "tier2", "P-99", "New document")
        with pytest.raises(Forbidden):
            portal.restore_builtin_template(who, "p-05")


# ------------------------------------------------------------------ versions
def test_version_numbers_major_minor_and_variants(portal):
    labels = [official(portal, change="major")["label"], official(portal, change="minor")["label"],
              official(portal, change="minor")["label"], official(portal, change="major")["label"],
              official(portal, variant="Product A", change="minor")["label"],
              official(portal, variant="  Product   A ", change="major")["label"]]
    assert labels == ["1.0", "1.1", "1.2", "2.0", "1.0", "2.0"]
    assert doclib.bump(["1.9", "2.0", "x"], "minor") == "2.1"
    with pytest.raises(Invalid):
        official(portal, change="huge")


def test_files_are_write_once_with_hash_and_listed_newest_first(portal):
    a = official(portal, comment="first")
    b = official(portal, change="minor")
    assert a["sha256"] == hashlib.sha256(DOCX).hexdigest() and a["size"] == len(DOCX) and a["status"] == "draft"
    assert a["evidence_run"] != b["evidence_run"]
    name, data = portal.doc_file(ADMIN, "p-05", a["RowKey"])
    assert name == "access-control.docx" and data == DOCX
    rows = portal.doc_versions(ADMIN, "p-05")["official"]
    assert [r["label"] for r in rows] == ["1.1", "1.0"]
    assert "doclib.upload" in actions(portal)


def test_file_list_includes_templates_for_admins_only(portal):
    files = portal.library_files(ADMIN, "tier2")
    assert len(files) == 13 and all(f["status"] == "template" and f["filename"].endswith(".docx") for f in files)
    assert portal.library_files(AUDITOR, "tier2")[0]["sha256"] and portal.library_files(CONTRIB) == []
    v = official(portal)
    portal.upload_template(ADMIN, "p-05", "mine.docx", DOCX, "")
    portal.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved")
    files = portal.library_files(ADMIN, "tier2")
    assert files[0]["status"] == "approved" and [f["filename"] for f in files if f["PartitionKey"] == "p-05"][-1] == "mine.docx"
    assert [f["status"] for f in portal.library_files(CONTRIB)] == ["approved"]


def test_status_flow_segregation_and_superseding(portal):
    v1 = official(portal)
    with pytest.raises(Forbidden):  # ADMIN uploaded it, ADMIN2 exists
        portal.set_version_status(ADMIN, "p-05", v1["RowKey"], "approved")
    with pytest.raises(Invalid):
        portal.set_version_status(ADMIN, "p-05", v1["RowKey"], "in_review")  # only through Send for approval
    done = portal.set_version_status(ADMIN2, "p-05", v1["RowKey"], "approved")
    assert done["approved_by"] == ADMIN2.email and done["approval_method"] == "direct"
    v2 = official(portal, change="minor")
    portal.set_version_status(ADMIN2, "p-05", v2["RowKey"], "approved")
    states = {v["label"]: v["status"] for v in portal.doc_versions(ADMIN, "p-05")["official"]}
    assert states == {"1.0": "approved", "1.1": "superseded"} or states == {"1.1": "approved", "1.0": "superseded"}
    assert states["1.1"] == "approved" and states["1.0"] == "superseded"
    # another product's approved version is not superseded
    other = official(portal, variant="Product A")
    portal.set_version_status(ADMIN2, "p-05", other["RowKey"], "approved")
    assert {c["variant"] for c in portal.library_doc(ADMIN, "p-05")["current"]} == {"", "Product A"}
    with pytest.raises(Invalid):
        portal.set_version_status(ADMIN2, "p-05", v1["RowKey"], "approved")  # superseded cannot be approved again
    with pytest.raises(Invalid):
        portal.set_version_status(ADMIN2, "p-05", v2["RowKey"], "archived")  # needs a reason
    portal.set_version_status(ADMIN2, "p-05", v2["RowKey"], "archived", "Withdrawn")
    assert "doclib.status.approved" in actions(portal)


def test_sole_admin_may_self_approve_and_it_is_marked(tmp_path):
    portal = make_portal(tmp_path, tenant="solo", people=(ADMIN, AUDITOR))
    v = portal.upload_official(ADMIN, "p-05", "a.docx", DOCX, "", "major", "")
    done = portal.set_version_status(ADMIN, "p-05", v["RowKey"], "approved")
    assert done["self_approved"] is True
    rec = [r for r in portal.log.records() if r["action"] == "doclib.status.approved"][-1]
    assert rec["detail"]["self_approved"] is True


def test_staff_download_only_approved_files(portal):
    v = official(portal)
    for who in (CONTRIB,):
        with pytest.raises(Invalid):
            portal.doc_file(who, "p-05", v["RowKey"])
    portal.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved")
    assert portal.doc_file(CONTRIB, "p-05", v["RowKey"])[1] == DOCX
    assert "doclib.download" in actions(portal)  # non-admin downloads are audited
    with pytest.raises(Invalid):
        portal.doc_file(CONTRIB, "p-05", "template")  # staff never get templates
    with pytest.raises(Invalid):
        portal.doc_file(ADMIN, "p-05", "../../etc/passwd")


# ------------------------------------------------------------------ templates
def test_template_replace_history_and_restore(portal):
    name, built_in = portal.doc_file(ADMIN, "p-05", "template")
    assert name == "p-05.docx" and built_in == doclib.builtin_template("p-05")[1]
    assert portal.library_doc(ADMIN, "p-05")["template_source"] == "Built-in"
    mine = built_in + b"extra"  # still starts with the zip signature
    portal.upload_template(ADMIN, "p-05", "my-template.docx", mine, "Company version")
    assert portal.library_doc(AUDITOR, "p-05")["template_source"] == "Custom"
    assert portal.doc_file(AUDITOR, "p-05", "template") == ("my-template.docx", mine)
    portal.restore_builtin_template(ADMIN, "p-05")
    assert portal.doc_file(ADMIN, "p-05", "template")[1] == built_in
    hist = portal.doc_versions(ADMIN, "p-05")["templates"]
    assert len(hist) == 2 and hist[0]["builtin"] and hist[1]["filename"] == "my-template.docx"
    assert portal.doc_file(ADMIN, "p-05", hist[1]["RowKey"])[1] == mine  # old template still downloadable
    assert {"doclib.template.replace", "doclib.template.restore"} <= set(actions(portal))


# ------------------------------------------------------------------ details and custom documents
def test_edit_title_keeps_number_and_is_audited(portal):
    d = portal.edit_doc(ADMIN, "p-05", {"title": "  Access   Control and Identity Policy ", "owner": HR.email,
                                        "classification": "Internal", "next_review": "2027-10-01"})
    assert d["title"] == "Access Control and Identity Policy" and d["number"] == "P-05"
    assert d["owner_name"] == "HR Lead" and d["owner_is_person"] and d["builtin_title"] == "Access Control Policy"
    rec = [r for r in portal.log.records() if r["action"] == "doclib.edit"][-1]
    assert rec["detail"]["before"]["title"] == "Access Control Policy" and rec["detail"]["after"]["next_review"] == "2027-10-01"
    n = len(portal.log.records())
    portal.edit_doc(ADMIN, "p-05", {"title": "Access Control and Identity Policy"})
    assert len(portal.log.records()) == n  # no change, no entry
    for bad in ({"title": "x"}, {"classification": "Secret"}, {"next_review": "01/02/2027"}):
        with pytest.raises(Invalid):
            portal.edit_doc(ADMIN, "p-05", bad)


def test_add_custom_document_and_review_dates(portal):
    d = portal.add_doc(ADMIN, "tier2", "P-14", "Privacy Policy", "policy")
    assert d["custom"] and d["id"].startswith("x-p-14-") and portal.library(ADMIN, folder="tier2")["docs"][-1]["id"] == d["id"]
    with pytest.raises(Invalid):
        portal.add_doc(ADMIN, "nope", "", "Some title")
    v = portal.upload_official(ADMIN, d["id"], "privacy.docx", DOCX, "", "major", "")
    portal.set_version_status(ADMIN2, d["id"], v["RowKey"], "approved")
    portal.edit_doc(ADMIN, d["id"], {"next_review": "2026-01-01"})
    assert portal.library_doc(ADMIN, d["id"])["review_state"] == "overdue"
    portal.edit_doc(ADMIN, d["id"], {"next_review": "2026-11-01"})
    assert portal.library_doc(ADMIN, d["id"])["review_state"] == "soon"
    with pytest.raises(Invalid):
        portal.doc_file(ADMIN, d["id"], "template")  # a custom document has no built-in template


# ------------------------------------------------------------------ upload safety and tenants
def test_upload_validation(portal):
    for name, data in (("run.exe", b"MZ\x90"), ("a.docm", DOCX), ("a.docx", b"not a zip"), ("a.docx", b""),
                       ("a.pdf", b"PK\x03\x04")):
        with pytest.raises(Invalid):
            portal.upload_official(ADMIN, "p-05", name, data, "", "major", "")
    with pytest.raises(Invalid):
        portal.upload_official(ADMIN, "no-such-doc", "a.docx", DOCX, "", "major", "")
    portal.upload_official(ADMIN, "p-05", "../../evil name.PDF", PDF, "", "major", "")
    assert portal.doc_versions(ADMIN, "p-05")["official"][0]["filename"] == "evil_name.PDF"


def test_other_tenant_sees_nothing(tmp_path):
    base = SqliteStore()
    one = make_portal(tmp_path, tenant="one", base_store=base)
    two = make_portal(tmp_path, tenant="two", base_store=base)
    v = one.upload_official(ADMIN, "p-05", "a.docx", DOCX, "", "major", "")
    one.edit_doc(ADMIN, "p-05", {"title": "Tenant one title"})
    assert two.doc_versions(ADMIN, "p-05")["official"] == [] and two.library_doc(ADMIN, "p-05")["title"] == "Access Control Policy"
    with pytest.raises(Invalid):
        two.doc_file(ADMIN, "p-05", v["RowKey"])
    with pytest.raises(Invalid):
        two.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved")


# ------------------------------------------------------------------ web pages
def test_pages_render_under_strict_csp(portal, client):
    v = official(portal, variant="Product A")
    portal.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved")
    portal.upload_template(ADMIN, "p-05", "t.docx", DOCX, "")
    paths = ["/library", "/library?f=tier2", "/library?view=files", "/library?q=access&view=files", "/library/p-05", "/library/isp",
             "/library/p-05-s1", "/library/soc2-03", "/library?f=other", "/library?view=register", "/library?edit=1",
             "/library/p-05?edit=1"]
    for who in (ADMIN, AUDITOR):
        as_(client, who)
        for path in paths:
            r = client.get(path)
            assert r.status_code == 200, (who.role, path, r.text[:300])
            assert 'style="' not in r.text and "<script>" not in r.text and "onclick" not in r.text
    as_(client, ADMIN)
    assert 'class="on">All files' in client.get("/library?view=files").text  # not the program-wide "view" variable
    assert 'class="on">Documents' in client.get("/library").text
    page = client.get("/library/p-05").text
    assert "Upload an official version" in page and "Upload template" in page and "Product A" in page
    as_(client, AUDITOR)
    page = client.get("/library/p-05").text
    assert "Upload an official version" not in page and "Upload template" not in page and "/library/p-05/files/template" in page


def test_staff_pages_show_only_published(portal, client):
    as_(client, CONTRIB)
    assert "No approved documents" in client.get("/library").text
    assert client.get("/library/p-05").status_code == 404
    v = official(portal)
    portal.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved")
    page = client.get("/library/p-05").text
    assert client.get("/library/p-05").status_code == 200 and "Download template" not in page and "Replace the template" not in page
    assert client.get(f"/library/p-05/files/{v['RowKey']}").status_code == 200
    assert client.get("/library/p-05/files/template").status_code == 404


def test_forms_need_csrf_and_role_and_download_is_attachment(portal, client):
    as_(client, ADMIN)
    token = csrf(client, "/library")
    files = {"file": ("ac.docx", DOCX, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    assert client.post("/library/p-05/official", data={"change": "major"}, files=files, follow_redirects=False).status_code == 403
    r = client.post("/library/p-05/official", data={"csrf": token, "change": "major", "variant": "Product A"}, files=files,
                    follow_redirects=False)
    assert r.status_code == 303 and "msg=" in r.headers["location"]
    vid = portal.doc_versions(ADMIN, "p-05")["official"][0]["RowKey"]
    r = client.get(f"/library/p-05/files/{vid}")
    assert r.headers["content-disposition"].startswith("attachment") and "sandbox" in r.headers["content-security-policy"]
    r = client.post(f"/library/p-05/versions/{vid}/status", data={"csrf": token, "status": "approved"}, follow_redirects=False)
    assert r.status_code == 403  # the uploader cannot approve while another admin exists
    assert portal.doc_versions(ADMIN, "p-05")["official"][0]["status"] == "draft"
    r = client.post("/library/p-05/edit", data={"csrf": token, "title": "Renamed Access Policy"}, follow_redirects=False)
    assert r.status_code == 303 and portal.library_doc(ADMIN, "p-05")["title"] == "Renamed Access Policy"
    r = client.post("/library/p-05/official", data={"csrf": token, "change": "major"}, follow_redirects=False)
    assert "err=" in r.headers["location"]  # no file chosen
    for who in (CONTRIB, AUDITOR):
        as_(client, who)
        tok = csrf(client, "/policies")
        assert client.post("/library/p-05/edit", data={"csrf": tok, "title": "Hacked title"}, follow_redirects=False).status_code == 403
        assert client.post("/library/add", data={"csrf": tok, "folder": "tier2", "title": "Hacked"}, follow_redirects=False).status_code == 403


def test_tables_are_in_backups(portal):
    from grcpal.portal import store

    assert {"docs", "docversions"} <= set(store.TABLES) and not {"docs", "docversions"} & set(store.NOT_BACKED_UP)
    assert zipfile.is_zipfile(io.BytesIO(DOCX))


# ------------------------------------------------------------------ batch 3: trash, restore, descriptions
def test_deleted_documents_go_to_trash_and_can_be_restored(portal):
    official(portal)  # p-05 has a file in the evidence store
    custom = portal.add_doc(ADMIN, "other", "X-1", "Old badge procedure")
    portal.trash_docs(ADMIN, ["p-05", custom["id"]], "Not used")
    shown = {d["id"] for d in portal.library(ADMIN)["docs"]}
    assert "p-05" not in shown and custom["id"] not in shown
    with pytest.raises(Invalid):
        portal.library_doc(ADMIN, "p-05")  # hidden everywhere, even by direct link
    trash = portal.library_trash(ADMIN)
    assert {d["id"] for d in trash["docs"]} == {"p-05", custom["id"]}
    assert portal.library_trash(AUDITOR)["docs"]  # auditors can see what was deleted
    portal.restore_doc(ADMIN, "p-05")
    d = portal.library_doc(ADMIN, "p-05")
    assert d["official_count"] == 1  # the file was never destroyed
    assert {"doclib.delete", "doclib.restore"} <= set(actions(portal))
    entry = next(r for r in portal.log.records() if r["action"] == "doclib.delete")
    assert entry["detail"]["reason"] == "Not used"


def test_only_admins_delete_or_restore_and_staff_never_see_the_trash(portal):
    for who in (CONTRIB, AUDITOR):
        with pytest.raises(Forbidden):
            portal.trash_docs(who, ["p-05"])
    portal.trash_docs(ADMIN, ["p-05"])
    with pytest.raises(Forbidden):
        portal.restore_doc(AUDITOR, "p-05")
    with pytest.raises(Forbidden):
        portal.library_trash(CONTRIB)
    with pytest.raises(Invalid):
        portal.trash_docs(ADMIN, ["p-05"])  # already in the Trash
    with pytest.raises(Invalid):
        portal.trash_docs(ADMIN, [])


def test_trash_is_tenant_scoped(tmp_path):
    shared = SqliteStore(tmp_path / "shared.db")
    a = make_portal(tmp_path / "a", tenant="acme", base_store=shared)
    b = make_portal(tmp_path / "b", tenant="globex", base_store=shared)
    a.trash_docs(ADMIN, ["p-05"])
    assert not b.library_trash(ADMIN)["docs"] and b.library_doc(ADMIN, "p-05")


def test_old_templates_can_be_removed_and_restored(portal):
    first = portal.upload_template(ADMIN, "p-05", "first.docx", DOCX, "")
    second = portal.upload_template(ADMIN, "p-05", "second.docx", DOCX, "")
    assert portal.doc_versions(ADMIN, "p-05")["active_template"]["RowKey"] == second["RowKey"]
    portal.remove_template(ADMIN, "p-05", second["RowKey"])
    v = portal.doc_versions(ADMIN, "p-05")
    assert v["active_template"]["RowKey"] == first["RowKey"]  # falls back to the next newest
    assert [t["filename"] for t in v["templates"]] == ["first.docx"]
    assert [t["filename"] for t in portal.library_trash(ADMIN)["templates"]] == ["second.docx"]
    portal.remove_template(ADMIN, "p-05", first["RowKey"])
    assert portal.doc_versions(ADMIN, "p-05")["active_template"] is None  # back to the built-in template
    portal.restore_template(ADMIN, "p-05", second["RowKey"])
    assert portal.doc_versions(ADMIN, "p-05")["active_template"]["RowKey"] == second["RowKey"]
    with pytest.raises(Forbidden):
        portal.remove_template(AUDITOR, "p-05", second["RowKey"])


def test_title_and_description_are_editable(portal):
    portal.edit_doc(ADMIN, "p-05", {"title": "Access control policy (2026)", "summary": "Who gets access and how."})
    d = portal.library_doc(ADMIN, "p-05")
    assert d["title"] == "Access control policy (2026)" and d["summary"] == "Who gets access and how."
    with pytest.raises(Invalid):
        portal.edit_doc(ADMIN, "p-05", {"summary": "x" * 1001})


def test_edit_mode_and_trash_pages(portal, client):
    as_(client, ADMIN)
    assert "Edit list" in client.get("/library").text
    page = client.get("/library?edit=1").text
    assert 'name="ids" value="p-05"' in page and "data-confirm=" in page
    r = client.post("/library/trash", data={"csrf": csrf(client, "/library"), "ids": ["p-05", "p-06"]},
                    follow_redirects=False)
    assert r.status_code == 303
    trash = client.get("/library/trash").text
    assert "Restore" in trash and client.get("/library/p-05").status_code == 404
    as_(client, CONTRIB)
    assert 'name="ids"' not in client.get("/library?edit=1").text  # edit mode is admins only
    assert client.get("/library/trash").status_code == 403
    r = client.post("/library/trash", data={"csrf": csrf(client, "/library"), "ids": ["p-07"]}, follow_redirects=False)
    assert r.status_code == 403


# ------------------------------------------------------------------ PR: one Details card, people, numbers, labels
def test_owner_and_approver_are_active_people(portal):
    d = portal.library_doc(ADMIN, "isp")
    assert not d["approver_is_person"] and d["approver_name"] == "Executive sponsor"  # built-in role until picked
    assert portal.edit_doc(ADMIN, "isp", {"approver": d["approver"]})["approver"] == d["approver"]  # unchanged role is kept
    d = portal.edit_doc(ADMIN, "isp", {"approver": ADMIN2.email, "owner": ADMIN.email})
    assert d["approver_is_person"] and d["approver_name"] == "VP" and d["owner_name"] == "Sec Lead"
    for bad in ({"approver": AUDITOR.email}, {"approver": "ceo@other.example"}, {"owner": "IT lead"}):
        with pytest.raises(Invalid):
            portal.edit_doc(ADMIN, "isp", bad)  # auditors, strangers and free text are refused
    assert portal.edit_doc(ADMIN, "isp", {"approver": ""})["approver"] == ""  # can be cleared
    with pytest.raises(Forbidden):
        portal.edit_doc(CONTRIB, "isp", {"approver": CONTRIB.email})


def test_document_number_is_editable_unique_and_audited(portal):
    v = official(portal)
    d = portal.edit_doc(ADMIN, "p-05", {"number": " POL-05 "})
    assert d["number"] == "POL-05" and d["id"] == "p-05"  # the internal id never changes
    assert portal.doc_versions(ADMIN, "p-05")["official"][0]["RowKey"] == v["RowKey"]  # versions still attached
    rec = [r for r in portal.log.records() if r["action"] == "doclib.edit"][-1]
    assert rec["detail"]["before"]["number"] == "P-05" and rec["detail"]["after"]["number"] == "POL-05"
    for bad in ("p-06", "P-06", "<b>x</b>", "x" * 21):
        with pytest.raises(Invalid):
            portal.edit_doc(ADMIN, "p-05", {"number": bad})  # taken (any case), unsafe or too long
    with pytest.raises(Invalid):
        portal.add_doc(ADMIN, "tier2", "pol-05", "Another policy")
    assert portal.edit_doc(ADMIN, "p-05", {"number": "P-05"})["number"] == "P-05"


def test_version_number_typed_on_upload_and_locked_after_approval(portal):
    v = portal.upload_official(ADMIN, "p-05", "a.docx", DOCX, "", "major", "", label="v3.2")
    assert v["label"] == "3.2"
    assert portal.doc_versions(ADMIN, "p-05")["next"] == {"major": "4.0", "minor": "3.3"}
    with pytest.raises(Invalid):
        portal.upload_official(ADMIN, "p-05", "b.docx", DOCX, "", "major", "", label="3.2")  # already used
    with pytest.raises(Invalid):
        portal.upload_official(ADMIN, "p-05", "b.docx", DOCX, "", "major", "", label="three")
    assert portal.upload_official(ADMIN, "p-05", "b.docx", DOCX, "Product A", "major", "", label="3.2")["label"] == "3.2"
    w = official(portal)  # left empty: numbered automatically
    assert w["label"] == "4.0"
    fixed = portal.set_version_label(ADMIN, "p-05", w["RowKey"], "3.3.1")
    assert fixed["label"] == "3.3.1"
    rec = [r for r in portal.log.records() if r["action"] == "doclib.version.label"][-1]
    assert rec["detail"]["from"] == "4.0" and rec["detail"]["to"] == "3.3.1" and rec["detail"]["sha256"] == w["sha256"]
    with pytest.raises(Invalid):
        portal.set_version_label(ADMIN, "p-05", w["RowKey"], "3.2")  # duplicate
    with pytest.raises(Forbidden):
        portal.set_version_label(AUDITOR, "p-05", w["RowKey"], "5.0")
    portal.set_version_status(ADMIN2, "p-05", w["RowKey"], "approved")
    with pytest.raises(Invalid):
        portal.set_version_label(ADMIN, "p-05", w["RowKey"], "5.0")  # approved: locked
    assert portal.store.get("docversions", "p-05", w["RowKey"])["label"] == "3.3.1"


def test_template_rename_keeps_file_and_extension(portal):
    t = portal.upload_template(ADMIN, "p-05", "old-name.docx", DOCX, "")
    r = portal.rename_template(ADMIN, "p-05", t["RowKey"], "Access Control Template 2026")
    assert r["filename"] == "Access_Control_Template_2026.docx" and r["original_filename"] == "old-name.docx"
    assert portal.doc_file(ADMIN, "p-05", "template") == ("Access_Control_Template_2026.docx", DOCX)  # same bytes
    rec = [r for r in portal.log.records() if r["action"] == "doclib.template.rename"][-1]
    assert rec["detail"]["sha256"] == t["sha256"] and rec["detail"]["from"] == "old-name.docx"
    for bad in ("other.pdf", "", "   "):
        with pytest.raises(Invalid):
            portal.rename_template(ADMIN, "p-05", t["RowKey"], bad)
    with pytest.raises(Forbidden):
        portal.rename_template(AUDITOR, "p-05", t["RowKey"], "x")
    portal.remove_template(ADMIN, "p-05", t["RowKey"])
    v = portal.doc_versions(ADMIN, "p-05")
    assert v["uploads"] == [] and [x["RowKey"] for x in v["removed"]] == [t["RowKey"]] and v["builtin"]["active"]
    with pytest.raises(Invalid):
        portal.rename_template(ADMIN, "p-05", t["RowKey"], "again")  # removed templates are not renamed


def test_details_card_edit_mode_and_compact_templates_page(portal, client):
    portal.upload_template(ADMIN, "p-05", "removed.docx", DOCX, "")
    gone = portal.doc_versions(ADMIN, "p-05")["uploads"][0]
    portal.remove_template(ADMIN, "p-05", gone["RowKey"])
    portal.upload_template(ADMIN, "p-05", "kept.docx", DOCX, "")
    draft = official(portal)
    as_(client, ADMIN)
    page = client.get("/library/p-05").text
    assert page.count("<h2>Details</h2>") == 1 and "Edit details" not in page and 'href="/library/p-05?edit=1#details"' in page
    assert "Choose file" not in page and 'data-autosubmit' in page and "kept.docx" in page and "removed.docx" in page
    assert 'aria-label="Rename kept.docx"' in page and 'aria-label="Remove kept.docx"' in page
    assert f'/versions/{draft["RowKey"]}/label' in page
    page = client.get("/library/p-05?edit=1").text
    for name in ("title", "number", "summary", "owner", "approver", "classification", "next_review"):
        assert f'name="{name}"' in page, name
    assert f'value="{ADMIN2.email}"' in page and f'value="{AUDITOR.email}"' not in page  # auditors are not offered
    token = csrf(client, "/library/p-05")
    r = client.post("/library/p-05/edit", data={"csrf": token, "title": "Access Control Policy", "number": "AC-1",
                                                "approver": ADMIN2.email, "owner": ADMIN.email}, follow_redirects=False)
    assert r.status_code == 303 and portal.library_doc(ADMIN, "p-05")["number"] == "AC-1"
    r = client.post(f"/library/p-05/versions/{draft['RowKey']}/label", data={"csrf": token, "label": "1.4"},
                    follow_redirects=False)
    assert r.status_code == 303 and portal.store.get("docversions", "p-05", draft["RowKey"])["label"] == "1.4"
    kept = portal.doc_versions(ADMIN, "p-05")["uploads"][0]
    r = client.post("/library/p-05/template", data={"csrf": token, "action": "rename", "version": kept["RowKey"],
                                                    "name": "Renamed"}, follow_redirects=False)
    assert r.status_code == 303 and portal.doc_versions(ADMIN, "p-05")["uploads"][0]["filename"] == "Renamed.docx"
    r = client.post("/library/p-05/official", data={"csrf": token, "change": "major", "label": "7.0"},
                    files={"file": ("x.docx", DOCX)}, follow_redirects=False)
    assert r.status_code == 303 and official(portal, change="minor")["label"] == "7.1"
    as_(client, AUDITOR)
    page = client.get("/library/p-05?edit=1").text
    assert 'name="number"' not in page and "?edit=1" not in page and 'aria-label="Rename' not in page  # read-only
    r = client.post(f"/library/p-05/versions/{draft['RowKey']}/label", data={"csrf": csrf(client, "/library/p-05"),
                                                                            "label": "9.0"}, follow_redirects=False)
    assert r.status_code == 403


# ------------------------------------------------------------------ PR 1b: list editing, register, reserved numbers
def test_numbers_of_trashed_documents_stay_reserved(portal):
    portal.trash_docs(ADMIN, ["p-06"])
    with pytest.raises(Invalid, match="in the Trash"):
        portal.edit_doc(ADMIN, "p-05", {"number": "P-06"})
    with pytest.raises(Invalid):
        portal.add_doc(ADMIN, "tier2", "p-06", "New policy")
    portal.restore_doc(ADMIN, "p-06")
    numbers = [d["number"] for d in portal.library(ADMIN)["docs"] if d["number"]]
    assert len(numbers) == len(set(n.lower() for n in numbers))  # no clash after restoring


def test_register_rows_export_and_access(portal):
    v = official(portal)
    portal.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved")
    official(portal, did="p-06")
    portal.edit_doc(ADMIN, "p-05", {"approver": ADMIN2.email})
    rows = {d["id"]: d for d in portal.document_register(AUDITOR)}
    assert len(rows) == len(portal.library(ADMIN)["docs"])
    assert rows["p-05"]["state"] == "Approved" and rows["p-05"]["versions"] == "v1.0" and rows["p-05"]["approver_name"] == "VP"
    assert rows["p-06"]["state"] == "Draft or in review" and rows["isp"]["state"] == "No version yet"
    data = portal.register_xlsx(AUDITOR)
    sheet = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert "Access Control Policy" in sheet and "Approver" in sheet and "VP" in sheet
    assert [r for r in portal.log.records() if r["action"] == "doclib.register.export"][-1]["actor"] == AUDITOR.email
    with pytest.raises(Forbidden):
        portal.document_register(CONTRIB)


def test_list_pencil_register_tab_and_export_route(portal, client):
    as_(client, ADMIN)
    page = client.get("/library?edit=1&f=tier2").text
    assert 'form="rn-p-05"' in page and 'id="rn-p-05"' in page and 'aria-label="Edit number and title of' in page
    assert page.count("<form") - page.count('<form id="rn-') >= 1  # row forms sit outside the delete form
    r = client.post("/library/p-05/edit", data={"csrf": csrf(client, "/library"), "number": "AC-05",
                                                "title": "Access Control Policy", "back": "list", "f": "tier2"},
                    follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/library?edit=1&f=tier2")
    assert portal.library_doc(ADMIN, "p-05")["number"] == "AC-05"
    r = client.post("/library/p-05/edit", data={"csrf": csrf(client, "/library"), "title": "Access Control Policy",
                                                "back": "list", "f": "//evil.example"}, follow_redirects=False)
    assert r.headers["location"].startswith("/library?edit=1") and "evil" not in r.headers["location"]
    page = client.get("/library?view=register").text
    assert "Export to Excel" in page and "AC-05" in page and 'class="on">Register' in page
    r = client.get("/library/register.xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"] and r.content[:2] == b"PK"
    as_(client, AUDITOR)
    assert client.get("/library/register.xlsx").status_code == 200
    assert 'form="rn-' not in client.get("/library?edit=1").text
    as_(client, CONTRIB)
    assert client.get("/library/register.xlsx").status_code == 403
    page = client.get("/library?view=register").text
    assert "Export to Excel" not in page and ">Register<" not in page  # staff fall back to the normal list


# ------------------------------------------------------------------ approval workflow
def mails(portal, to=None):
    return [m for m in portal.mailer.outbox if to is None or m.to == to]


def sent_for_approval(portal, approver=HR, did="p-05", **kw):
    portal.edit_doc(ADMIN, did, {"approver": approver.email})
    v = official(portal, did=did, **kw)
    return portal.request_approval(ADMIN, did, v["RowKey"], "Please review before Friday")


def test_send_for_approval_emails_only_the_named_approver(portal):
    v = official(portal)
    with pytest.raises(Invalid, match="approver"):
        portal.request_approval(ADMIN, "p-05", v["RowKey"])  # built-in role, not a person yet
    portal.edit_doc(ADMIN, "p-05", {"approver": ADMIN.email})
    with pytest.raises(Invalid, match="cannot be"):
        portal.request_approval(ADMIN, "p-05", v["RowKey"])  # the uploader cannot approve their own upload
    portal.edit_doc(ADMIN, "p-05", {"approver": HR.email})
    with pytest.raises(Forbidden):
        portal.request_approval(CONTRIB, "p-05", v["RowKey"])
    w = portal.request_approval(ADMIN, "p-05", v["RowKey"], "Please review")
    assert w["status"] == "in_review" and w["approval_approver"] == HR.email and w["approval_requested_by"] == ADMIN.email
    m = mails(portal, HR.email)[0]
    assert "Approval needed" in m.subject and f"/library/p-05/review/{v['RowKey']}" in m.text and len(mails(portal)) == 1
    assert "access-control.docx" not in m.text  # a link, never the file
    assert portal.pending_approvals(HR)["mine"][0]["RowKey"] == v["RowKey"]
    assert portal.pending_approvals(ADMIN)["waiting"][0]["approver_name"] == "HR Lead"
    assert portal.pending_approvals(CONTRIB) == {"mine": [], "waiting": []}
    with pytest.raises(Invalid):
        portal.request_approval(ADMIN, "p-05", v["RowKey"])  # already waiting
    assert "doclib.approval.request" in actions(portal)


def test_approver_approves_after_reading_and_it_is_evidence(portal):
    old = official(portal)
    portal.set_version_status(ADMIN2, "p-05", old["RowKey"], "approved")
    v = sent_for_approval(portal)
    assert portal.library_doc(HR, "p-05")  # the approver can open the document while it waits for them
    assert portal.doc_file(HR, "p-05", v["RowKey"])[1] == DOCX
    with pytest.raises(Invalid, match="read"):
        portal.decide_approval(HR, "p-05", v["RowKey"], "approve", "", read=False)
    done = portal.decide_approval(HR, "p-05", v["RowKey"], "approve", "OK for 2027", read=True)
    assert done["status"] == "approved" and done["approval_method"] == "workflow" and done["read_confirmed"]
    assert done["approved_by"] == HR.email and done["decision_note"] == "OK for 2027"
    assert portal.store.get("docversions", "p-05", old["RowKey"])["status"] == "superseded"
    rec = [r for r in portal.log.records() if r["action"] == "doclib.status.approved"][-1]
    assert rec["detail"]["method"] == "workflow" and rec["detail"]["sha256"] == v["sha256"]
    assert rec["detail"]["uploaded_by"] == ADMIN.email and rec["actor"] == HR.email
    run = portal.evidence.list_runs("doclib_approval")[-1]
    record = __import__("json").loads(portal.evidence.read("doclib_approval", run, "decision.json"))
    assert record["decision"] == "approved" and record["sha256"] == v["sha256"] and record["by"] == HR.email
    assert any("Approved:" in m.subject for m in mails(portal, ADMIN.email))
    assert portal.pending_approvals(HR)["mine"] == []  # the My work item closes by itself
    with pytest.raises(Invalid):
        portal.decide_approval(HR, "p-05", v["RowKey"], "approve", "", read=True)  # no double approval


def test_deny_needs_a_reason_and_returns_to_draft(portal):
    v = sent_for_approval(portal)
    with pytest.raises(Invalid, match="why"):
        portal.decide_approval(HR, "p-05", v["RowKey"], "deny", "  ")
    done = portal.decide_approval(HR, "p-05", v["RowKey"], "deny", "Section 4 needs the backup owner")
    assert done["status"] == "draft" and done["denied_by"] == HR.email
    m = [m for m in mails(portal, ADMIN.email) if "Not approved" in m.subject][0]
    assert "Section 4 needs the backup owner" in m.text
    assert "doclib.approval.deny" in actions(portal) and portal.library(HR)["docs"] == []  # nothing published
    with pytest.raises(Invalid):
        portal.library_doc(HR, "p-05")  # access ended with the request
    again = portal.request_approval(ADMIN, "p-05", v["RowKey"])  # can be sent again after a fix note
    assert again["status"] == "in_review"


def test_only_the_named_approver_decides_and_sees_nothing_else(portal):
    v = sent_for_approval(portal)
    other = official(portal, did="p-06")  # an unrelated draft
    for who in (CONTRIB, AUDITOR, ADMIN2, ADMIN):
        with pytest.raises(Forbidden):
            portal.decide_approval(who, "p-05", v["RowKey"], "approve", "", read=True)
    with pytest.raises(Invalid):
        portal.library_doc(CONTRIB, "p-05")
    with pytest.raises(Invalid):
        portal.doc_file(HR, "p-06", other["RowKey"])  # the approver sees only the version sent to them
    with pytest.raises(Invalid):
        portal.approval_view(CONTRIB, "p-05", v["RowKey"])
    assert [x["RowKey"] for x in portal.doc_versions(HR, "p-05")["official"]] == [v["RowKey"]]
    assert portal.doc_versions(HR, "p-05")["templates"] == []
    with pytest.raises(Invalid):
        portal.request_approval(ADMIN, "p-06", other["RowKey"])  # p-06 has no person as approver yet


def test_withdraw_reassign_and_direct_approval_close_the_request(portal):
    v = sent_for_approval(portal)
    with pytest.raises(Invalid):
        portal.reassign_approval(ADMIN, "p-05", v["RowKey"], ADMIN.email)  # the uploader
    w = portal.reassign_approval(ADMIN, "p-05", v["RowKey"], CONTRIB.email)
    assert w["approval_approver"] == CONTRIB.email
    assert any("No longer needed" in m.subject for m in mails(portal, HR.email))
    assert any("Approval needed" in m.subject for m in mails(portal, CONTRIB.email))
    with pytest.raises(Forbidden):
        portal.decide_approval(HR, "p-05", v["RowKey"], "approve", "", read=True)  # no longer theirs
    portal.withdraw_approval(ADMIN, "p-05", v["RowKey"], "Wrong file")
    assert portal.store.get("docversions", "p-05", v["RowKey"])["status"] == "draft"
    assert portal.pending_approvals(CONTRIB)["mine"] == []
    portal.request_approval(ADMIN, "p-05", v["RowKey"])  # back to HR, the document's approver
    done = portal.set_version_status(ADMIN2, "p-05", v["RowKey"], "approved", "Signed copy inside the document")
    assert done["approval_method"] == "direct" and done["approved_by"] == ADMIN2.email
    assert portal.pending_approvals(HR)["mine"] == []
    assert sum("No longer needed" in m.subject for m in mails(portal, HR.email)) == 2
    assert {"doclib.approval.reassign", "doclib.approval.withdraw"} <= set(actions(portal))


def test_reminders_every_seven_days(portal):
    import datetime as dt

    v = sent_for_approval(portal)
    assert portal.send_approval_reminders(ADMIN2) == []  # not due yet
    portal._today = dt.date.today() + dt.timedelta(days=8)
    sent = portal.send_approval_reminders(ADMIN2)
    assert [r["version"] for r in sent] == [v["label"]] and any("Reminder" in m.subject for m in mails(portal, HR.email))
    assert portal.send_approval_reminders(ADMIN2) == []  # once per 7 days
    portal.decide_approval(HR, "p-05", v["RowKey"], "approve", "", read=True)
    portal._today = dt.date.today() + dt.timedelta(days=30)
    assert portal.send_approval_reminders(ADMIN2) == []  # nothing waiting any more
    assert "doclib.approval.remind" in actions(portal)


def test_review_page_my_work_and_forms(portal, client):
    v = sent_for_approval(portal)
    as_(client, HR)
    work = client.get("/work").text
    assert f"/library/p-05/review/{v['RowKey']}" in work and "Waiting for your approval" in work
    page = client.get(f"/library/p-05/review/{v['RowKey']}").text
    assert 'name="decision" value="approve"' in page and 'name="read"' in page
    assert 'style="' not in page and "<script>" not in page and "onclick" not in page
    token = csrf(client, "/work")
    r = client.post(f"/library/p-05/review/{v['RowKey']}", data={"csrf": token, "decision": "approve"}, follow_redirects=False)
    assert r.status_code == 303 and "err=" in r.headers["location"]  # read box not ticked
    r = client.post(f"/library/p-05/review/{v['RowKey']}", data={"csrf": token, "decision": "approve", "read": "1"},
                    follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/work")
    assert portal.store.get("docversions", "p-05", v["RowKey"])["status"] == "approved"
    assert "This request is closed" in client.get(f"/library/p-05/review/{v['RowKey']}").text
    w = official(portal, change="minor")
    portal.request_approval(ADMIN, "p-05", w["RowKey"])
    as_(client, CONTRIB)
    assert client.get(f"/library/p-05/review/{w['RowKey']}").status_code == 404
    r = client.post(f"/library/p-05/review/{w['RowKey']}", data={"csrf": csrf(client, "/work"), "decision": "approve",
                                                                 "read": "1"}, follow_redirects=False)
    assert r.status_code == 403
    as_(client, ADMIN)
    page = client.get("/library/p-05").text
    assert "Waiting for <b>hr@example.com</b>" in page and "Reassign" in page and "Withdraw" in page
    assert "Documents waiting on others" in client.get("/work").text
    as_(client, AUDITOR)
    page = client.get(f"/library/p-05/review/{w['RowKey']}").text
    assert "waiting for" in page and 'name="decision"' not in page
