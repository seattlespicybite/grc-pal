"""Build the Document library templates (Word files) and the catalog that lists them.

Run from the repository root after editing tools/doctemplates/*.yaml:

    python tools/build_doc_templates.py

Needs python-docx (development only; the running portal just serves the finished files, so there is no
runtime dependency). Output: grcpal/starter/doctemplates/files/<id>.docx and catalog.yaml.
"""
from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "tools" / "doctemplates"
OUT = ROOT / "grcpal" / "starter" / "doctemplates"

# Colors follow the portal (static/app.css): dark teal brand, teal accent, light teal surface.
BRAND, ACCENT, SURFACE, BORDER, MUTED = "0E2A35", "0F8B8D", "EEF3F4", "DDE5E8", "5C6F78"
FONT = "Calibri"
ROLE_NAMES = {"security-lead": "Security lead", "hr-lead": "HR lead", "it-lead": "IT lead",
              "engineering-lead": "Engineering lead", "legal-lead": "Legal lead",
              "executive-sponsor": "Executive sponsor"}
KINDS = {"policy": "Policy", "standard": "Standard", "procedure": "Procedure / plan", "soc2": "SOC 2 document"}


def role(r: str) -> str:
    return ROLE_NAMES.get(r, r)


# ------------------------------------------------------------------ low-level helpers
def shade(cell, fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"), el.set(qn("w:color"), "auto"), el.set(qn("w:fill"), fill)
    pr.append(el)


def borders(table, color: str = BORDER) -> None:
    tbl = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single"), b.set(qn("w:sz"), "4"), b.set(qn("w:space"), "0"), b.set(qn("w:color"), color)
        el.append(b)
    tbl.append(el)


def repeat_header(row) -> None:
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def widths(table, cm: list[float]) -> None:
    from docx.shared import Cm

    table.autofit = False
    for col, w in zip(table.columns, cm):
        col.width = Cm(w)
    for row in table.rows:
        for cell, w in zip(row.cells, cm):
            cell.width = Cm(w)


def run(par, text: str, *, bold=False, italic=False, color: str | None = None, size: float | None = None):
    r = par.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
    return r


def cell_text(cell, text: str, *, bold=False, color: str | None = None, italic=False, size: float = 10) -> None:
    cell.text = ""
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    run(par, text, bold=bold, color=color, italic=italic, size=size)


def field(par, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r._r.append(f)
        else:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            r._r.append(t)


def new_document(d: dict) -> Document:
    doc = Document()
    sec = doc.sections[0]
    from docx.shared import Cm

    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in (("Title", 24, BRAND), ("Heading 1", 15, BRAND), ("Heading 2", 12, ACCENT)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
        st.font.color.rgb = RGBColor.from_string(color)
        st.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    props = doc.core_properties
    props.title, props.author, props.subject = d["title"], "[Company Name]", f"{d['number']} template"
    props.created = props.modified = dt.datetime(2026, 1, 1)
    head = sec.header.paragraphs[0]
    run(head, f"[Company Name]  ·  {d['number']}  ·  {d['title']}", color=MUTED, size=8)
    foot = sec.footer.paragraphs[0]
    run(foot, "Internal  ·  Version [1.0]  ·  Page ", color=MUTED, size=8)
    field(foot, "PAGE")
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return doc


def title_block(doc, d: dict) -> None:
    run(doc.add_paragraph(), KINDS[d["kind"]].upper() + " TEMPLATE", bold=True, color=ACCENT, size=9)
    doc.add_paragraph(d["title"], style="Title")
    p = doc.add_paragraph()
    run(p, "[Company Name]", bold=True, color=BRAND)
    run(p, f"  ·  {d['number']}", color=MUTED)


def notice(doc, text: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    borders(t, ACCENT)
    shade(t.rows[0].cells[0], SURFACE)
    cell_text(t.rows[0].cells[0], text, color=BRAND, size=9.5)
    doc.add_paragraph()


def control_table(doc, d: dict, reg: dict) -> None:
    parent = reg.get(d.get("parent", ""))
    rows = [
        ("Document number", d["number"]), ("Title", d["title"]), ("Type", KINDS[d["kind"]]),
        ("Owner", f"[Named person in the role: {role(d['owner'])}]" if d.get("owner") else "[Named owner: one role, not a team]"),
        ("Approver", role(d["approver"]) if d.get("approver") else "[Document owner and Security lead]"),
        ("Version", "[1.0]  (major on any change to a commitment; minor for wording)"),
        ("Effective date", "[YYYY-MM-DD]"), ("Next review date", "[Within 12 months of approval]"),
        ("Classification", "Internal  (mark customer-shareable versions as such)"),
        ("Parent document", f"{parent['number']} {parent['title']}" if parent else ("Information Security Policy" if d["kind"] != "policy" or d["id"] != "isp" else "None (top level)")),
        ("Framework mapping", " · ".join(x for x in (f"SOC 2: {d['soc2']}" if d.get("soc2") else "",
                                                      f"ISO 27001:2022: {d['iso']}" if d.get("iso") else "") if x) or "[SOC 2 criteria and other framework references]"),
        ("GRC-pal controls", d.get("controls") or "[Control IDs that prove this document]"),
    ]
    t = doc.add_table(rows=len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for r, (k, v) in zip(t.rows, rows):
        shade(r.cells[0], SURFACE)
        cell_text(r.cells[0], k, bold=True, color=BRAND)
        cell_text(r.cells[1], v)
    widths(t, [4.2, 12.4])
    doc.add_paragraph()


def table(doc, head: list[str], rows: list[list[str]], cm: list[float]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    borders(t)
    for cell, text in zip(t.rows[0].cells, head):
        shade(cell, BRAND)
        cell_text(cell, text, bold=True, color="FFFFFF", size=9.5)
    repeat_header(t.rows[0])
    for r, vals in zip(t.rows[1:], rows):
        for cell, text in zip(r.cells, vals):
            cell_text(cell, text)
    widths(t, cm)
    doc.add_paragraph()


def guide(doc, text: str) -> None:
    run(doc.add_paragraph(), "Guidance: " + text, italic=True, color=MUTED, size=9)


def tail(doc, d: dict, extra_related: str = "") -> None:
    doc.add_heading("Exceptions", level=1)
    doc.add_paragraph("Exceptions are requested in writing with the reason, the risk, compensating controls and an end date, "
                      "approved by [role] before they take effect, and recorded in the policy exception register.")
    if d["kind"] == "policy":
        doc.add_heading("Compliance and enforcement", level=1)
        doc.add_paragraph("Violations of this policy may lead to disciplinary action up to and including termination, "
                          "or contract termination for third parties, as set out in the Human Resources Security Policy.")
    doc.add_heading("Related documents", level=1)
    doc.add_paragraph(extra_related or "[Parent policy, child standards and procedures, and records this document produces]")
    doc.add_heading("Revision history", level=1)
    table(doc, ["Date", "Version", "Author", "Summary of change"], [["[YYYY-MM-DD]", "[1.0]", "[Name]", "Initial version"]],
          [3.0, 2.2, 3.8, 7.6])


# ------------------------------------------------------------------ builders per kind
def build_policy(d: dict, c: dict, reg: dict):
    doc = new_document(d)
    title_block(doc, d)
    notice(doc, "TEMPLATE. Replace every [bracketed] text with what [Company Name] actually does, delete the grey "
                "guidance lines, and remove statements that are not true today. Every \"must\" is something an auditor can "
                "test, so keep the \"How it is tested\" column accurate. Set the document owner, approver and review dates in "
                "the document control table.")
    control_table(doc, d, reg)
    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(c["purpose"].strip())
    doc.add_heading("2. Scope and applicability", level=1)
    doc.add_paragraph(c["scope"].strip())
    n = 3
    for sec in c["sections"]:
        doc.add_heading(f"{n}. {sec['h']}", level=1)
        if sec.get("guide"):
            guide(doc, sec["guide"])
        rows = []
        for i, item in enumerate(sec["items"], 1):
            text, _, test = item.partition(" || ")
            rows.append([f"{n}.{i}", text, test])
        table(doc, ["Ref", "Statement", "How it is tested"], rows, [1.3, 9.6, 5.7])
        n += 1
    doc.add_heading(f"{n}. Roles and responsibilities", level=1)
    table(doc, ["Role", "Responsibility"], [r.split(" || ") for r in c["roles"]], [4.6, 12.0])
    tail(doc, d, c.get("related", ""))
    if d["id"] != "isp":
        doc.add_heading("Acknowledgment", level=1)
        doc.add_paragraph("Personnel acknowledge this policy at hire and at least annually. Acknowledgments are recorded per version.")
    return doc


def build_standard(d: dict, reg: dict):
    doc = new_document(d)
    title_block(doc, d)
    parent = reg[d["parent"]]
    notice(doc, "TEMPLATE (skeleton). A standard holds the measurable values its parent policy points to, so values can change "
                "without re-approving the policy. Decide each value from current practice, then record why. Delete the grey guidance lines.")
    control_table(doc, d, reg)
    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(f"This standard sets the parameters for the {parent['title']} ({parent['number']}).")
    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph("[Systems, data, people and locations the parameters apply to.]")
    doc.add_heading("3. Parameters", level=1)
    guide(doc, "State each value as something that can be measured and evidenced, for example a number of days or a setting name.")
    table(doc, ["Parameter", "Value", "Why this value / source", "Evidence"],
          [[p, "[value]", "[decision basis]", "[report or setting that proves it]"] for p in d["params"]], [4.6, 3.2, 4.8, 4.0])
    doc.add_heading("4. Applicability and exceptions", level=1)
    doc.add_paragraph("[Systems or groups that follow a different value, and the approved exception route.]")
    tail(doc, d, f"{parent['number']} {parent['title']}")
    return doc


def build_procedure(d: dict, reg: dict):
    doc = new_document(d)
    title_block(doc, d)
    parent = reg[d["parent"]] if d.get("parent") in reg else None
    notice(doc, "TEMPLATE (skeleton). A procedure describes the steps people actually perform and names the record each step "
                "leaves; this is what the auditor walks through and samples. Write what happens today, not what you hope will "
                "happen. Delete the grey guidance lines.")
    control_table(doc, d, reg)
    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(f"This document describes how [Company Name] carries out the requirements of "
                      f"{(parent['number'] + ' ' + parent['title']) if parent else 'its parent policy'}.")
    doc.add_heading("2. Scope, trigger and frequency", level=1)
    doc.add_paragraph("[What starts this procedure (event, ticket or schedule), how often it runs, and what it covers.]")
    doc.add_heading("3. Roles", level=1)
    table(doc, ["Role", "Responsibility"], [["[Role]", "[What this role does in the procedure]"]], [4.6, 12.0])
    doc.add_heading("4. Steps", level=1)
    guide(doc, "One row per action. Name the system used and the record produced so each step can be sampled.")
    table(doc, ["Step", "Who", "Action", "Record left behind"],
          [[str(i), "[Role]", "[Action]", "[Ticket, log or report]"] for i in (1, 2, 3)], [1.3, 3.0, 7.6, 4.7])
    doc.add_heading("5. Records produced", level=1)
    doc.add_paragraph(f"Expected records: {d['record']}. State where they are kept and for how long.")
    doc.add_heading("6. Measures", level=1)
    doc.add_paragraph("[How completion and timeliness are checked, who reviews them and how often.]")
    tail(doc, d, f"{parent['number']} {parent['title']}" if parent else "")
    return doc


def build_soc2(d: dict, reg: dict):
    doc = new_document(d)
    title_block(doc, d)
    notice(doc, "TEMPLATE (skeleton). This document is required by SOC 2 and is written by management, not by the auditor. "
                "Complete each section from what the company really does and keep it current. Delete the grey guidance lines.")
    control_table(doc, d, reg)
    doc.add_heading("Why this document exists", level=1)
    doc.add_paragraph(d["why"])
    doc.add_heading("1. Purpose and scope", level=1)
    doc.add_paragraph("[What this document covers, the services and the period it applies to.]")
    doc.add_heading("2. Content", level=1)
    guide(doc, "Add the sections your auditor and customers need. Reference controls and evidence rather than repeating them.")
    table(doc, ["Item", "Description", "Owner", "Evidence / reference"], [["[Item]", "[Description]", "[Role]", "[Reference]"]],
          [3.6, 6.6, 2.6, 3.8])
    doc.add_heading("3. Review", level=1)
    doc.add_paragraph("[Who reviews this document, how often, and what triggers an earlier review.]")
    tail(doc, d)
    return doc


def main() -> int:
    reg_file = yaml.safe_load((SRC / "registry.yaml").read_text())
    content = yaml.safe_load((SRC / "policies.yaml").read_text())
    docs = reg_file["docs"]
    reg = {d["id"]: d for d in docs}
    (OUT / "files").mkdir(parents=True, exist_ok=True)
    for old in (OUT / "files").glob("*.docx"):
        old.unlink()
    catalog_docs = []
    for d in docs:
        kind = d["kind"]
        if kind == "policy":
            doc = build_policy(d, content[d["id"]], reg)
        elif kind == "standard":
            doc = build_standard(d, reg)
        elif kind == "procedure":
            doc = build_procedure(d, reg)
        else:
            doc = build_soc2(d, reg)
        name = f"{d['id']}.docx"
        doc.save(OUT / "files" / name)
        entry = {k: v for k, v in d.items() if k not in ("params", "why", "record", "folds")}
        entry["file"] = name
        for k in ("folds", "record", "why"):
            if d.get(k):
                entry["summary"] = d[k]
        if d.get("params"):
            entry["summary"] = "Parameters: " + "; ".join(d["params"])
        catalog_docs.append(entry)
    out = {"folders": reg_file["folders"], "docs": catalog_docs}
    (OUT / "catalog.yaml").write_text("# Generated by tools/build_doc_templates.py. Do not edit by hand.\n"
                                      + yaml.safe_dump(out, sort_keys=False, allow_unicode=True, width=120))
    print(f"built {len(docs)} templates")
    return 0


if __name__ == "__main__":
    sys.exit(main())
